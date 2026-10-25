"""对资源使用导出逐表汇总；不猜测跨表读者映射，不导出身份或论文全文。"""
import csv
import hashlib
import json
import warnings
from collections import Counter
from datetime import date, datetime, timezone
from pathlib import Path

SOURCES = {
    "seats": ("座位预约.csv", 0, "学号", "预约日期", ["楼层", "状态"]),
    "rooms": ("研讨室预约.xlsx", 1, "账号", "日期", ["研讨室"]),
    "gate_events": ("门禁刷卡.xlsx", 0, "ReadNum", "Cdate", ["InPort"]),
    "gate_counts": ("门禁入馆人次.xlsx", 0, "ReadNum", None, []),
    "readers": ("读者数据.csv", 0, "ID", None, ["TYPE"]),
    "scholars": ("学者库数据.xlsx", 0, None, "发表日期", ["论文分类"]),
}


def text(value):
    return "" if value is None else str(value).strip()


def rows(path, header_row):
    if path.suffix.lower() == ".csv":
        try:
            # 验证整文件编码，避免读到中间才发现需换编码。
            path.read_text(encoding="utf-8-sig")
            encoding = "utf-8-sig"
        except UnicodeDecodeError:
            encoding = "gb18030"
        with path.open(encoding=encoding, newline="") as stream:
            reader = csv.DictReader(stream)
            if reader.fieldnames is None or len(reader.fieldnames) != len(set(reader.fieldnames)):
                raise ValueError("CSV 表头缺失或重复。")
            for row in reader:
                if None in row or any(value is None for value in row.values()):
                    raise ValueError("CSV 列数不一致。")
                yield row
    else:
        from openpyxl import load_workbook
        with warnings.catch_warnings():
            warnings.filterwarnings("ignore", message="Cannot parse header or footer.*", category=UserWarning)
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                iterator = workbook.worksheets[0].iter_rows(values_only=True)
                for _ in range(header_row):
                    next(iterator)
                headers = [text(value) for value in next(iterator)]
                if not all(headers) or len(headers) != len(set(headers)):
                    raise ValueError("Excel 表头缺失或重复。")
                for values in iterator:
                    if any(value is not None for value in values):
                        yield dict(zip(headers, values))
            finally:
                workbook.close()


def observed_date(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()[:10]
    value = text(value)
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y%m%d", "%Y-%m-%d %H:%M:%S", "%Y/%m/%d %H:%M:%S", "%Y-%m-%d %H:%M:%S.%f"):
        try:
            return datetime.strptime(value, fmt).date().isoformat()
        except ValueError:
            pass
    # 年/月精度保留为原精度，不能补成该月第一天或当年1月1日。
    for fmt, output in (("%Y-%m", "%Y-%m"), ("%Y", "%Y")):
        try:
            return datetime.strptime(value, fmt).strftime(output)
        except ValueError:
            pass
    return None


def summarize(path, config):
    _, header_row, reader_field, date_field, dimensions = config
    readers, seen, distributions, days = set(), set(), {key: Counter() for key in dimensions}, Counter()
    count = missing_readers = duplicates = invalid_dates = invalid_counts = 0
    date_precisions = Counter()
    gate_total = 0
    for row in rows(path, header_row):
        required = set(dimensions) | ({reader_field} if reader_field else set()) | ({date_field} if date_field else set())
        if not required.issubset(row):
            raise ValueError("来源表头与预期不符。")
        count += 1
        # 仅为诊断统计重复，仍保留所有行；摘要集合不输出。
        digest = hashlib.sha256(json.dumps({key: text(value) for key, value in row.items()}, sort_keys=True, ensure_ascii=False).encode()).digest()
        duplicates += digest in seen
        seen.add(digest)
        if reader_field:
            reader = text(row[reader_field])
            if reader:
                readers.add(reader)
            else:
                missing_readers += 1
        for key in dimensions:
            distributions[key][text(row[key]) or "未提供"] += 1
        if date_field:
            day = observed_date(row[date_field])
            if day:
                days[day] += 1
                date_precisions[{4: "year", 7: "month", 10: "day"}[len(day)]] += 1
            else:
                invalid_dates += 1
        if "CNT" in row:
            value = text(row["CNT"])
            if value.isdigit():
                gate_total += int(value)
            else:
                invalid_counts += 1
    return {"source_file": path.name, "source_sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "rows": count, "distinct_source_readers": len(readers) if reader_field else None,
            "missing_reader_rows": missing_readers if reader_field else None,
            "duplicate_excess_rows": duplicates, "invalid_date_rows": invalid_dates if date_field else None,
            "date_precision_counts": dict(date_precisions),
            "observed_date_min": min(days) if days else None, "observed_date_max": max(days) if days else None,
            "reported_gate_count_sum": gate_total if "gate" in path.name or path.name == "门禁入馆人次.xlsx" else None,
            "invalid_count_rows": invalid_counts,
            "dimensions": {key: [{"name": name, "count": n} for name, n in sorted(counter.items(), key=lambda item: (-item[1], item[0])) if n >= 5]
                           for key, counter in distributions.items()},
            "suppressed_dimension_rows": {key: sum(n for n in counter.values() if n < 5) for key, counter in distributions.items()}}


def analyze_usage(directory, dataset_label="unknown", progress=None):
    if dataset_label not in {"unknown", "simulated", "real"}:
        raise ValueError("数据标签无效。")
    if not Path(directory).is_dir():
        raise ValueError("输入资料目录不存在。")
    datasets = {}
    for key, config in SOURCES.items():
        path = Path(directory) / config[0]
        if path.exists():
            datasets[key] = summarize(path, config)
            if progress:
                progress(key + ": " + str(datasets[key]["rows"]) + " rows")
    return {"dataset_label": dataset_label, "built_at": datetime.now(timezone.utc).isoformat(), "is_live": False,
            "datasets": datasets, "limitations": ["各表单独统计，不推断不同脱敏读者或书目编号的关联。",
                "记录数不等于利用率；缺少开放容量、时段和统一统计周期，不能比较利用率。",
                "重复流水只报告不删除，日期范围仅为当前导出可解析记录的观察范围。",
                "部分发表日期仅精确到年或月，保留原精度；不支持的日期文本单独计数，不补造日期。",
                "分组计数少于5的项不输出名称，抑制的记录数单独计数。"]}
