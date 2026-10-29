import asyncio
from io import BytesIO

import pytest
from fastapi import HTTPException
from openpyxl import Workbook

import app.api.data_management as data_management_api
import app.services.data_manager as data_manager_module
from app.agent.agent import agent
from app.config import settings
from app.services.data_manager import data_manager


def admin_headers():
    return {"X-Admin-Token": settings.ADMIN_TOKEN}


def test_data_routes_require_admin(client):
    assert client.get("/api/admin/data/sources").status_code == 401
    assert client.get("/api/admin/data/files").status_code == 401
    assert client.post("/api/admin/data/quick-sync-all").status_code == 401
    assert client.post(
        "/api/admin/data/upload/reader",
        files={"file": ("reader.csv", b"id,name\n1,reader\n", "text/csv")},
    ).status_code == 401


def test_upload_rejects_path_components():
    with pytest.raises(HTTPException) as exc:
        data_management_api._validate_filename("../outside.csv")
    assert exc.value.status_code == 400


def test_upload_size_limit_cleans_partial_file(client, monkeypatch):
    monkeypatch.setattr(data_management_api, "MAX_UPLOAD_BYTES", 8)
    response = client.post(
        "/api/admin/data/upload/reader",
        files={"file": ("too-large.csv", b"id,name\n1,reader\n", "text/csv")},
        headers=admin_headers(),
    )
    assert response.status_code == 413
    assert not (data_management_api.IMPORTED_DIR / "reader" / "too-large.csv").exists()


def test_xlsx_uncompressed_size_limit_rejects_upload_and_cleans_file(client, monkeypatch):
    monkeypatch.setattr(data_manager_module, "MAX_XLSX_UNCOMPRESSED_BYTES", 1)
    workbook = Workbook()
    workbook.active.append(["header"])
    workbook.active.append(["value"])
    buffer = BytesIO()
    workbook.save(buffer)

    response = client.post(
        "/api/admin/data/upload/exam",
        files={"file": ("expanded.xlsx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=admin_headers(),
    )
    assert response.status_code == 422
    assert not (data_management_api.IMPORTED_DIR / "exam" / "expanded.xlsx").exists()


def test_csv_row_limit_rejects_import_without_manifest_update(monkeypatch):
    monkeypatch.setattr(data_manager_module, "MAX_IMPORT_ROWS", 1)
    category_dir = data_manager_module.IMPORTED_DIR / "reader"
    category_dir.mkdir(parents=True, exist_ok=True)
    source = category_dir / "too-many.csv"
    source.write_text("id,name\n1,first\n2,second\n", encoding="utf-8")

    with pytest.raises(ValueError, match="行数超过"):
        data_manager.process_file("reader", source)
    assert not data_manager.manifest_file.exists()


def test_xlsx_row_limit_rejects_import_without_manifest_update(monkeypatch):
    monkeypatch.setattr(data_manager_module, "MAX_IMPORT_ROWS", 1)
    category_dir = data_manager_module.IMPORTED_DIR / "exam"
    category_dir.mkdir(parents=True, exist_ok=True)
    source = category_dir / "too-many.xlsx"
    workbook = Workbook()
    workbook.active.append(["header"])
    workbook.active.append(["first"])
    workbook.active.append(["second"])
    workbook.save(source)

    with pytest.raises(ValueError, match="行数超过"):
        data_manager.process_file("exam", source)
    assert not data_manager.manifest_file.exists()


def test_upload_does_not_overwrite_existing_file(client):
    path = data_management_api.IMPORTED_DIR / "reader" / "same-name.csv"
    content = b"id,name\n1,first\n"
    first = client.post(
        "/api/admin/data/upload/reader",
        files={"file": ("same-name.csv", content, "text/csv")},
        headers=admin_headers(),
    )
    second = client.post(
        "/api/admin/data/upload/reader",
        files={"file": ("same-name.csv", b"id,name\n2,second\n", "text/csv")},
        headers=admin_headers(),
    )
    assert first.status_code == 200
    assert second.status_code == 409
    assert path.read_bytes() == content


def test_failed_parse_removes_uploaded_file_and_hides_internal_error(client, monkeypatch):
    def fail_parse(_category, _path):
        raise RuntimeError("internal path /data/private/file.csv")

    monkeypatch.setattr(data_manager, "process_file", fail_parse)
    response = client.post(
        "/api/admin/data/upload/reader",
        files={"file": ("invalid.csv", b"id,name\n1,reader\n", "text/csv")},
        headers=admin_headers(),
    )
    assert response.status_code == 422
    assert "internal path" not in response.text
    assert not (data_management_api.IMPORTED_DIR / "reader" / "invalid.csv").exists()


def test_quick_sync_uses_configured_source_directory(client, tmp_path, monkeypatch):
    source_file = tmp_path / "raw" / "备考联盟数据" / "备考联盟.xlsx"
    source_file.parent.mkdir(parents=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["备考类型"])
    sheet.append(["考研"])
    workbook.save(source_file)
    monkeypatch.setattr(settings, "DATA_SOURCE_DIRECTORY", str(tmp_path / "raw"))

    response = client.post(
        "/api/admin/data/quick-sync/exam",
        headers=admin_headers(),
    )
    assert response.status_code == 200
    body = response.json()
    assert body["code"] == 0
    assert body["data"]["row_count"] == 1
    assert body["data"]["source_file"] == "备考联盟.xlsx"


def test_multisource_answers_do_not_claim_unimported_traffic(client):
    response = data_manager.query_multimodal_knowledge("图书馆现在人多吗")
    assert response["available"] is False
    assert "没有已导入" in response["reply"]
    assert "每日客流规律" not in response["reply"]
    tool, _args, payload, _cards, reply = agent._prepare("图书馆现在人多吗", reader_id=None)
    assert tool == "query_multimodal_data"
    assert payload["available"] is False
    assert "没有已导入" in reply


def test_imported_seat_summary_is_static_and_source_labeled(tmp_path):
    category_dir = data_manager_module.IMPORTED_DIR / "seats"
    category_dir.mkdir(parents=True, exist_ok=True)
    source = category_dir / "座位样本.csv"
    source.write_text("区域,预约\n二层,记录1\n三层,记录2\n", encoding="utf-8")
    data_manager.process_file("seats", source)

    response = data_manager.query_multimodal_knowledge("自习区现在有空位吗")
    assert response["available"] is True
    assert response["row_count"] == 2
    assert response["source_file"] == "座位样本.csv"
    assert "静态历史样本" in response["reply"]
    assert "不能说明当前空位" in response["reply"]


def test_multisource_tool_bypasses_llm_rewrite(monkeypatch):
    fallback = "仅根据导入数据给出的固定回答。"
    monkeypatch.setattr(agent, "_llm_ready", lambda: True)
    result = asyncio.run(
        agent._compose_reply(
            "问题", "query_multimodal_data", {}, {}, fallback, []
        )
    )
    assert result == fallback
